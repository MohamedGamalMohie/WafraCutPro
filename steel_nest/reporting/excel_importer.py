"""
Excel and CSV Import Module for SteelNest Pro
Supports robust parsing, header detection, column mapping, and error reporting.
"""
import csv
import uuid
from typing import List, Tuple, Dict, Any, Optional
import openpyxl

from steel_nest.models.part import PartCutItem, PlatePartItem
from steel_nest.models.stock import StockItem
from steel_nest.config import MaterialType

class ExcelDataImporter:
    @staticmethod
    def import_unified_cut_list(file_path: str) -> Tuple[List[PartCutItem], List[PlatePartItem], List[str]]:
        """
        Imports both 1D linear parts and 2D plate parts from an Excel or CSV file.
        Automatically categorizes rows into linear (width = 0) and plates (width > 0).
        Returns (linear_parts, plate_parts, error_messages).
        """
        linear_parts: List[PartCutItem] = []
        plate_parts: List[PlatePartItem] = []
        errors: List[str] = []

        if file_path.lower().endswith(".csv"):
            all_sheets_rows = {"Sheet1": ExcelDataImporter._read_csv_rows(file_path)}
        else:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                all_sheets_rows = {}
                for sheetname in wb.sheetnames:
                    if "تعليمات" in sheetname.lower() or "instruction" in sheetname.lower():
                        continue
                    ws = wb[sheetname]
                    rows = [list(r) for r in ws.iter_rows(values_only=True)]
                    if rows:
                        all_sheets_rows[sheetname] = rows
            except Exception as e:
                return [], [], [f"Error reading file: {str(e)}"]

        if not all_sheets_rows:
            return [], [], ["File is empty or could not be read."]

        target_map = {
            "mark": ["label", "mark", "part", "part mark", "part_mark", "tag", "item", "كود"],
            "designation": ["material", "mat", "section", "designation", "profile", "section designation", "size", "plate", "قطاع", "خامة"],
            "grade": ["grade", "steel grade", "material", "mat", "رتبة"],
            "length": ["length", "length ▲", "cut length", "len", "length (mm)", "cut length (mm)", "طول"],
            "width": ["width", "cut width", "width (mm)", "cut width (mm)", "عرض", "w", "w (mm)"],
            "quantity": ["quantity", "qty", "count", "pcs", "pieces", "عدد", "كمية"],
            "rotation": ["rotation", "allow rotation", "allow 90° rotation", "rot", "دوران"],
            "weight": ["weight", "wt", "weight (kg)", "وزن"],
            "notes": ["notes", "remark", "remarks", "comment", "description", "ملاحظات"]
        }

        for sheet_title, rows in all_sheets_rows.items():
            header_row_idx = None
            col_map = {}

            for r_idx in range(min(10, len(rows))):
                row = [str(c).strip().lower() if c is not None else "" for c in rows[r_idx]]
                matched = {}
                for key, aliases in target_map.items():
                    for c_idx, col_name in enumerate(row):
                        if any(alias == col_name or alias in col_name for alias in aliases):
                            matched[key] = c_idx
                            break
                if ("designation" in matched or "mark" in matched) and "length" in matched:
                    header_row_idx = r_idx
                    col_map = matched
                    break
                if "thickness" in matched and "length" in matched:
                    header_row_idx = r_idx
                    col_map = matched
                    break

            if header_row_idx is None:
                continue

            for r_idx in range(header_row_idx + 1, len(rows)):
                row = rows[r_idx]
                if not row or all(str(c).strip() == "" for c in row):
                    continue

                try:
                    mark = str(row[col_map["mark"]]).strip() if "mark" in col_map and col_map["mark"] < len(row) and row[col_map["mark"]] is not None else f"P-{r_idx}"
                    desig = str(row[col_map["designation"]]).strip() if "designation" in col_map and col_map["designation"] < len(row) and row[col_map["designation"]] is not None else ""
                    grade = str(row[col_map["grade"]]).strip() if "grade" in col_map and col_map["grade"] < len(row) and row[col_map["grade"]] is not None else "S355JR"

                    raw_len = row[col_map["length"]] if "length" in col_map and col_map["length"] < len(row) else None
                    if raw_len is None or str(raw_len).strip() == "":
                        continue
                    length_val = float(str(raw_len).replace(",", "").strip())

                    raw_width = row[col_map["width"]] if "width" in col_map and col_map["width"] < len(row) else 0.0
                    width_val = 0.0
                    if raw_width is not None and str(raw_width).strip() != "":
                        try:
                            width_val = float(str(raw_width).replace(",", "").strip())
                        except ValueError:
                            width_val = 0.0

                    raw_qty = row[col_map["quantity"]] if "quantity" in col_map and col_map["quantity"] < len(row) else 1
                    qty_val = 1
                    if raw_qty is not None and str(raw_qty).strip() != "":
                        try:
                            qty_val = int(float(str(raw_qty).replace(",", "").strip()))
                        except ValueError:
                            qty_val = 1

                    if length_val <= 0 or qty_val <= 0:
                        continue

                    rot = True
                    if "rotation" in col_map and col_map["rotation"] < len(row) and row[col_map["rotation"]] is not None:
                        rot_str = str(row[col_map["rotation"]]).strip().upper()
                        rot = rot_str not in ("NO", "FALSE", "0", "N", "لا")

                    note = str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) and row[col_map["notes"]] else ""

                    if width_val > 0:
                        thk = 10.0
                        desig_clean = desig.upper().replace("PL", "").replace("MM", "").replace("P", "").strip()
                        try:
                            thk = float(desig_clean)
                        except ValueError:
                            thk = 10.0
                        plate_parts.append(PlatePartItem(
                            id=str(uuid.uuid4()),
                            mark=mark,
                            thickness_mm=thk,
                            grade=grade,
                            width_mm=width_val,
                            length_mm=length_val,
                            quantity=qty_val,
                            allow_rotation=rot,
                            notes=note
                        ))
                    else:
                        wt = 0.0
                        if "weight" in col_map and col_map["weight"] < len(row) and row[col_map["weight"]] is not None:
                            try:
                                wt = float(str(row[col_map["weight"]]).replace(",", "").strip())
                            except ValueError:
                                wt = 0.0
                        linear_parts.append(PartCutItem(
                            id=str(uuid.uuid4()),
                            mark=mark,
                            designation=desig if desig else "IPE 300",
                            grade=grade,
                            length_mm=length_val,
                            quantity=qty_val,
                            weight_kg=wt,
                            priority=1,
                            notes=note
                        ))
                except Exception as e:
                    errors.append(f"Sheet '{sheet_title}', Row {r_idx + 1}: Parse error - {str(e)}")

        if not linear_parts and not plate_parts:
            lin, l_err = ExcelDataImporter.import_linear_cut_list(file_path)
            if lin:
                return lin, [], l_err
            pl, p_err = ExcelDataImporter.import_plate_cut_list(file_path)
            if pl:
                return [], pl, p_err
            errors.extend(l_err)

        return linear_parts, plate_parts, errors

    @staticmethod
    def import_linear_cut_list(file_path: str) -> Tuple[List[PartCutItem], List[str]]:
        """
        Imports linear cut pieces from .xlsx or .csv.
        Returns (list_of_parts, error_messages).
        """
        parts: List[PartCutItem] = []
        errors: List[str] = []

        if file_path.lower().endswith(".csv"):
            rows = ExcelDataImporter._read_csv_rows(file_path)
        else:
            rows = ExcelDataImporter._read_excel_rows(file_path)

        if not rows:
            return [], ["File is empty or could not be read."]

        # Header detection
        header_row_idx, col_map = ExcelDataImporter._find_headers(rows, {
            "mark": ["mark", "part", "part mark", "part_mark", "tag", "item"],
            "designation": ["section", "designation", "profile", "section designation", "size"],
            "grade": ["grade", "steel grade", "material", "mat"],
            "length": ["length", "cut length", "len", "length (mm)", "cut length (mm)"],
            "quantity": ["quantity", "qty", "count", "pcs", "pieces"],
            "weight": ["weight", "wt", "weight (kg)", "weight per piece (kg)"],
            "priority": ["priority", "prio"],
            "notes": ["notes", "remark", "remarks", "comment", "description"]
        })

        if header_row_idx is None:
            return [], ["Could not identify table headers in file. Please ensure columns include Mark, Section, Grade, Length, and Quantity."]

        for r_idx in range(header_row_idx + 1, len(rows)):
            row = rows[r_idx]
            if not row or all(str(c).strip() == "" for c in row):
                continue

            try:
                mark = str(row[col_map["mark"]]).strip() if "mark" in col_map and col_map["mark"] < len(row) else f"P-{r_idx}"
                desig = str(row[col_map["designation"]]).strip() if "designation" in col_map and col_map["designation"] < len(row) else ""
                grade = str(row[col_map["grade"]]).strip() if "grade" in col_map and col_map["grade"] < len(row) else "S355JR"
                
                raw_len = row[col_map["length"]] if "length" in col_map and col_map["length"] < len(row) else None
                raw_qty = row[col_map["quantity"]] if "quantity" in col_map and col_map["quantity"] < len(row) else 1
                
                if not mark or not desig or raw_len is None:
                    continue

                length_val = float(str(raw_len).replace(",", "").strip())
                qty_val = int(float(str(raw_qty).replace(",", "").strip()))

                wt = 0.0
                if "weight" in col_map and col_map["weight"] < len(row) and row[col_map["weight"]] is not None:
                    try:
                        wt = float(str(row[col_map["weight"]]).replace(",", "").strip())
                    except ValueError:
                        pass

                prio = 1
                if "priority" in col_map and col_map["priority"] < len(row) and row[col_map["priority"]] is not None:
                    try:
                        prio = int(float(str(row[col_map["priority"]]).strip()))
                    except ValueError:
                        pass

                note = str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) and row[col_map["notes"]] else ""

                if length_val <= 0 or qty_val <= 0:
                    errors.append(f"Row {r_idx + 1}: Skipped due to non-positive length ({length_val}) or quantity ({qty_val}).")
                    continue

                part = PartCutItem(
                    id=str(uuid.uuid4()),
                    mark=mark,
                    designation=desig,
                    grade=grade,
                    length_mm=length_val,
                    quantity=qty_val,
                    weight_kg=wt,
                    priority=prio,
                    notes=note
                )
                parts.append(part)
            except Exception as e:
                errors.append(f"Row {r_idx + 1}: Parse error - {str(e)}")

        return parts, errors

    @staticmethod
    def import_plate_cut_list(file_path: str) -> Tuple[List[PlatePartItem], List[str]]:
        """Imports 2D plate parts from .xlsx or .csv."""
        parts: List[PlatePartItem] = []
        errors: List[str] = []

        rows = ExcelDataImporter._read_csv_rows(file_path) if file_path.lower().endswith(".csv") else ExcelDataImporter._read_excel_rows(file_path)
        if not rows:
            return [], ["File is empty."]

        header_row_idx, col_map = ExcelDataImporter._find_headers(rows, {
            "mark": ["mark", "part", "part mark", "part_mark"],
            "thickness": ["thickness", "thk", "thickness (mm)", "thickness_mm"],
            "grade": ["grade", "steel grade", "material"],
            "width": ["width", "width (mm)", "w", "w (mm)"],
            "length": ["length", "len", "length (mm)", "l", "l (mm)"],
            "quantity": ["quantity", "qty", "count"],
            "rotation": ["rotation", "allow rotation", "allow 90° rotation", "rot"],
            "notes": ["notes", "remarks", "comment"]
        })

        if header_row_idx is None:
            return [], ["Could not identify plate headers. Ensure columns for Mark, Thickness, Grade, Width, Length, Quantity exist."]

        for r_idx in range(header_row_idx + 1, len(rows)):
            row = rows[r_idx]
            if not row or all(str(c).strip() == "" for c in row):
                continue
            try:
                mark = str(row[col_map["mark"]]).strip() if "mark" in col_map else f"PL-{r_idx}"
                thk = float(str(row[col_map["thickness"]]).strip())
                grade = str(row[col_map["grade"]]).strip() if "grade" in col_map else "S275JR"
                w = float(str(row[col_map["width"]]).strip())
                l = float(str(row[col_map["length"]]).strip())
                qty = int(float(str(row[col_map["quantity"]]).strip())) if "quantity" in col_map else 1
                
                rot = True
                if "rotation" in col_map and col_map["rotation"] < len(row) and row[col_map["rotation"]] is not None:
                    rot_str = str(row[col_map["rotation"]]).strip().upper()
                    rot = rot_str not in ("NO", "FALSE", "0", "N")

                notes = str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) else ""

                if thk <= 0 or w <= 0 or l <= 0 or qty <= 0:
                    errors.append(f"Row {r_idx + 1}: Invalid dimensions.")
                    continue

                parts.append(PlatePartItem(
                    id=str(uuid.uuid4()),
                    mark=mark,
                    thickness_mm=thk,
                    grade=grade,
                    width_mm=w,
                    length_mm=l,
                    quantity=qty,
                    allow_rotation=rot,
                    notes=notes
                ))
            except Exception as e:
                errors.append(f"Row {r_idx + 1}: Parse error - {str(e)}")

        return parts, errors

    @staticmethod
    def import_stock_inventory(file_path: str) -> Tuple[List[StockItem], List[str]]:
        """Imports stock inventory from .xlsx or .csv."""
        stock: List[StockItem] = []
        errors: List[str] = []

        if file_path.lower().endswith(".csv"):
            rows = ExcelDataImporter._read_csv_rows(file_path)
        else:
            # Check if there is a dedicated Stock sheet
            wb = openpyxl.load_workbook(file_path, data_only=True)
            target_sheet = None
            for sname in wb.sheetnames:
                sname_lower = sname.lower()
                if "stock" in sname_lower or "مخزن" in sname_lower or "خامات" in sname_lower:
                    target_sheet = sname
                    break
            rows = ExcelDataImporter._read_excel_rows(file_path, sheet_name=target_sheet)

        if not rows:
            return [], ["File is empty."]

        header_row_idx, col_map = ExcelDataImporter._find_headers(rows, {
            "id": ["label", "stock id", "id", "stock_id", "tag"],
            "type": ["type", "material type", "material_type"],
            "designation": ["material", "mat", "section", "designation", "profile"],
            "grade": ["grade", "steel grade", "material", "mat"],
            "length": ["length", "length ▲", "stock length", "length (mm)"],
            "width": ["width", "stock width", "width (mm)"],
            "thickness": ["thickness", "thickness (mm)"],
            "quantity": ["quantity", "available qty", "qty", "count"],
            "cost": ["price", "cost", "unit cost", "unit cost ($)"],
            "location": ["location", "bay", "rack", "yard"],
            "notes": ["notes", "remarks", "comment"]
        })

        if header_row_idx is None:
            return [], ["Could not identify headers for Stock Inventory."]

        for r_idx in range(header_row_idx + 1, len(rows)):
            row = rows[r_idx]
            if not row or all(str(c).strip() == "" for c in row):
                continue
            try:
                sid = str(row[col_map["id"]]).strip() if "id" in col_map and row[col_map["id"]] else f"STK-{uuid.uuid4().hex[:6]}"
                width_val = float(str(row[col_map["width"]]).strip()) if "width" in col_map and col_map["width"] < len(row) and row[col_map["width"]] else 0.0
                
                mtype_str = str(row[col_map["type"]]).strip().upper() if "type" in col_map and row[col_map["type"]] else ""
                if "PLATE" in mtype_str or width_val > 0:
                    mtype = MaterialType.PLATE
                else:
                    mtype = MaterialType.LINEAR
                
                desig = str(row[col_map["designation"]]).strip() if "designation" in col_map else ""
                grade = str(row[col_map["grade"]]).strip() if "grade" in col_map else "S355JR"
                if not desig and grade:
                    desig = grade
                length_val = float(str(row[col_map["length"]]).strip()) if "length" in col_map else 0.0
                thk_val = float(str(row[col_map["thickness"]]).strip()) if "thickness" in col_map and col_map["thickness"] < len(row) and row[col_map["thickness"]] else 0.0
                
                qty = int(float(str(row[col_map["quantity"]]).strip())) if "quantity" in col_map else 1
                cost = float(str(row[col_map["cost"]]).strip()) if "cost" in col_map and col_map["cost"] < len(row) and row[col_map["cost"]] else 0.0
                loc = str(row[col_map["location"]]).strip() if "location" in col_map and col_map["location"] < len(row) and row[col_map["location"]] else "Warehouse A"
                notes = str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) and row[col_map["notes"]] else ""

                if length_val <= 0 or qty <= 0:
                    errors.append(f"Row {r_idx + 1}: Non-positive stock length or qty.")
                    continue

                stock.append(StockItem(
                    id=sid,
                    designation=desig,
                    grade=grade,
                    length_mm=length_val,
                    width_mm=width_val,
                    thickness_mm=thk_val,
                    available_qty=qty,
                    used_qty=0,
                    cost_per_unit=cost,
                    is_remnant=False,
                    location=loc,
                    material_type=mtype,
                    notes=notes
                ))
            except Exception as e:
                errors.append(f"Row {r_idx + 1}: Parse error - {str(e)}")

        return stock, errors

    @staticmethod
    def _read_excel_rows(file_path: str, sheet_name: Optional[str] = None) -> List[List[Any]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
        return rows

    @staticmethod
    def _read_csv_rows(file_path: str) -> List[List[Any]]:
        rows = []
        with open(file_path, mode='r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.reader(f)
            for r in reader:
                rows.append(r)
        return rows

    @staticmethod
    def _find_headers(rows: List[List[Any]], target_map: Dict[str, List[str]]) -> Tuple[Optional[int], Dict[str, int]]:
        import re
        for r_idx in range(min(10, len(rows))):
            row = [str(c).strip().lower() if c is not None else "" for c in rows[r_idx]]
            matched = {}
            for key, aliases in target_map.items():
                for c_idx, col_name in enumerate(row):
                    if not col_name:
                        continue
                    # First check exact match
                    is_match = False
                    for alias in aliases:
                        a = alias.lower()
                        if a == col_name:
                            is_match = True
                            break
                        if len(a) > 2 and re.search(r'\b' + re.escape(a) + r'\b', col_name):
                            is_match = True
                            break
                        if len(a) > 3 and a in col_name:
                            is_match = True
                            break
                    if is_match:
                        matched[key] = c_idx
                        break
            # If we matched the primary essentials (e.g. designation & length, or mark)
            if "designation" in matched and "length" in matched:
                return r_idx, matched
            if "thickness" in matched and "width" in matched and "length" in matched:
                return r_idx, matched
            if "mark" in matched and "length" in matched:
                return r_idx, matched
            if "id" in matched and "length" in matched:
                return r_idx, matched
        return None, {}
